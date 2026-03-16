import json, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList

STATIC_DIR = os.path.join(os.path.dirname(__file__), 'static')

CATEGORY_MAP = {
    'Daimi İT xidməti': ['İT İnfrastruktur', 'Məlumat Bazası', 'Video, Monitorinq'],
    'İnfosec': ['İnformasiya Təhlükəsizlik', 'Təlim və Maarifləndirmə'],
    'Əlavə IT xidməti': ['HelpDesk və Texniki Dəstək'],
    'SAAS xidməti': ['Bulud Xidmətləri'],
    'ERP': ['Avtomatlaşdırılmış Sistemlər', 'SaaS Biznes Process'],
    'GRC': ['Audit və Uyğunluq'],
    'Layihə': ['Konsaltinq və Layihə'],
}

BOARD_CATS = ['Daimi İT xidməti', 'İnfosec', 'Əlavə IT xidməti', 'SAAS xidməti', 'ERP', 'GRC', 'Layihə']

GROUP_ORDER = ['Tabia', 'AFI', 'Azmade', 'PMD', 'Novex', 'Azersheker', 'Separated']

COLORS = {
    'header_bg': 'FF1B2A4A',
    'header_font': 'FFFFFFFF',
    'group_bg': 'FF2D4A7A',
    'group_font': 'FFFFFFFF',
    'alt_row': 'FFF0F4FA',
    'total_bg': 'FF1B2A4A',
    'total_font': 'FFFFFFFF',
    'border': 'FFB0BEC5',
    'kpi_bg': 'FFE8F5E9',
    'kpi_val': 'FF1B5E20',
}

MONTHS_AZ = ['Yanvar', 'Fevral', 'Mart', 'Aprel', 'May', 'İyun',
             'İyul', 'Avqust', 'Sentyabr', 'Oktyabr', 'Noyabr', 'Dekabr']


def load_data(adjustments=None):
    with open(os.path.join(STATIC_DIR, 'pricing_data.json'), 'r', encoding='utf-8') as f:
        data = json.load(f)
    legal_path = os.path.join(STATIC_DIR, 'company_legal_names.json')
    legal = {}
    if os.path.exists(legal_path):
        with open(legal_path, 'r', encoding='utf-8') as f:
            legal = json.load(f)
    return data, legal


def _cat_total(val):
    """Extract total from category value (handles both old flat and new nested format)."""
    if isinstance(val, dict):
        return val.get('total', 0) or 0
    return val or 0


def aggregate_board_cats(categories):
    """Map our 11 internal categories to the 7 board categories."""
    result = {bc: 0 for bc in BOARD_CATS}
    for board_cat, our_cats in CATEGORY_MAP.items():
        for c in our_cats:
            result[board_cat] += _cat_total(categories.get(c, 0))
    return result


def apply_adjustments(data, adj):
    if not adj:
        return data
    g = adj.get('global', 0) / 100
    groups = {k: v / 100 for k, v in adj.get('groups', {}).items()}
    cats = {k: v / 100 for k, v in adj.get('categories', {}).items()}
    comps = {k: v / 100 for k, v in adj.get('companies', {}).items()}
    result = {}
    for name, info in data.items():
        new_cats = {}
        multiplier_base = (1 + g) * (1 + groups.get(info['group'], 0)) * (1 + comps.get(name, 0))
        for cat, val in info['categories'].items():
            base = _cat_total(val)
            mult = multiplier_base * (1 + cats.get(cat, 0))
            adjusted = base * mult
            if isinstance(val, dict):
                # Scale individual service totals proportionally
                adj_services = []
                for svc in val.get("services", []):
                    adj_svc = dict(svc)
                    adj_svc['total'] = round((svc.get('total', 0) or 0) * mult, 2)
                    adj_svc['price'] = round((svc.get('price', 0) or 0) * mult, 2)
                    adj_services.append(adj_svc)
                new_cats[cat] = {"total": round(adjusted, 2), "services": adj_services}
            else:
                new_cats[cat] = round(adjusted, 2)
        monthly = sum(_cat_total(v) for v in new_cats.values())
        result[name] = {
            'group': info['group'],
            'categories': new_cats,
            'monthly': round(monthly, 2),
            'annual': round(monthly * 12, 2),
        }
    return result


def _parse_effective_month(effective_date):
    """Parse effective_date string (YYYY-MM-DD) and return the 0-based month index (0=Jan).
    Returns None if no date or invalid."""
    if not effective_date:
        return None
    try:
        parts = effective_date.split('-')
        return int(parts[1]) - 1  # 0-based: Jan=0, Jul=6, Dec=11
    except (IndexError, ValueError):
        return None


def _get_company_eff_month(company_name, group_name, adjustments, global_eff_date, category_name=None):
    """Get effective month for a specific company, checking hierarchy:
    company_date > category_date > group_date > global effective_date > default January (0).
    Returns 0-based month index."""
    if adjustments:
        # 1. Per-company date (highest priority)
        cd = (adjustments.get('company_dates') or {}).get(company_name)
        if cd:
            return _parse_effective_month(cd)
        # 2. Per-category date (if category specified)
        if category_name:
            cat_d = (adjustments.get('category_dates') or {}).get(category_name)
            if cat_d:
                return _parse_effective_month(cat_d)
        # 3. Per-group date
        gd = (adjustments.get('group_dates') or {}).get(group_name)
        if gd:
            return _parse_effective_month(gd)
    # 4. Global effective date
    m = _parse_effective_month(global_eff_date)
    if m is not None:
        return m
    # 5. Default: January (start of year)
    return 0


def add_monthly_sales_sheet(wb, base_data, adjusted_data, legal, effective_date=None, view='company', adjustments=None):
    """Add a sheet showing monthly sales breakdown.

    base_data = original data (no adjustments)
    adjusted_data = data with adjustments applied
    effective_date = 'YYYY-MM-DD' — global fallback date
    adjustments = original adjustments dict with group_dates/company_dates for per-group scheduling
    """
    global_eff_month = _parse_effective_month(effective_date)
    has_per_dates = adjustments and (adjustments.get('group_dates') or adjustments.get('company_dates'))

    thin = Side(style='thin', color=COLORS['border'])
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    hdr_fill = PatternFill('solid', fgColor=COLORS['header_bg'])
    hdr_font = Font(name='Arial', size=11, bold=True, color=COLORS['header_font'])
    grp_fill = PatternFill('solid', fgColor=COLORS['group_bg'])
    grp_font = Font(name='Arial', size=11, bold=True, color=COLORS['group_font'])
    tot_fill = PatternFill('solid', fgColor=COLORS['total_bg'])
    tot_font = Font(name='Arial', size=11, bold=True, color=COLORS['total_font'])
    normal = Font(name='Arial', size=10)
    bold = Font(name='Arial', size=10, bold=True)
    adj_font = Font(name='Arial', size=10, color='006600')  # green for adjusted months
    num_fmt = '#,##0.00 [$₼-42C]'
    now = datetime.datetime.now()
    adj_fill = PatternFill('solid', fgColor='FFE8F5E9')  # light green bg for adjusted months

    if view == 'company':
        ws = wb.create_sheet('Aylıq Satış (Şirkət)')
        ws.sheet_properties.tabColor = '7A4A2D'

        ws.merge_cells(f'A1:{get_column_letter(14)}1')
        title = f'Şirkət üzrə aylıq satış hesabatı — {now.year}'
        if effective_date:
            title += f'  (qiymət dəyişikliyi: {effective_date})'
        ws['A1'] = title
        ws['A1'].font = Font(name='Arial', size=14, bold=True, color='1B2A4A')

        # Headers
        headers = ['Şirkət'] + MONTHS_AZ + ['İllik Cəmi']
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=ci, value=h)
            c.font = hdr_font
            c.fill = hdr_fill
            c.border = border
            c.alignment = Alignment(horizontal='center', wrap_text=True)

        row = 4
        grand_total_rows = []

        for gi, group in enumerate(GROUP_ORDER):
            companies = [(n, adjusted_data[n]) for n in adjusted_data if adjusted_data[n]['group'] == group]
            if not companies:
                continue
            companies.sort(key=lambda x: x[1]['monthly'], reverse=True)

            for cc in range(1, len(headers) + 1):
                c = ws.cell(row=row, column=cc)
                c.fill = grp_fill
                c.font = grp_font
                c.border = border
            ws.cell(row=row, column=1, value=f'▸ {group}')
            row += 1
            grp_start = row

            for ci, (name, adj_info) in enumerate(companies):
                full_name = legal.get(name, name + ' MMC')
                ws.cell(row=row, column=1, value=full_name).font = normal
                base_monthly = base_data[name].get('monthly_total', base_data[name].get('monthly', 0)) if name in base_data else adj_info['monthly']
                adj_monthly = adj_info['monthly']
                # Per-company effective month (company > group > global)
                eff_month = _get_company_eff_month(name, group, adjustments, effective_date)

                for mi in range(12):
                    use_adjusted = (eff_month is None) or (mi >= eff_month)
                    if use_adjusted:
                        val = adj_monthly
                        cell_font = adj_font if eff_month is not None else normal
                    else:
                        val = base_monthly
                        cell_font = normal
                    c = ws.cell(row=row, column=2 + mi, value=val)
                    c.font = cell_font
                    c.number_format = num_fmt
                    if eff_month is not None and mi >= eff_month:
                        c.fill = adj_fill

                c = ws.cell(row=row, column=14, value=f'=SUM(B{row}:M{row})')
                c.font = bold
                c.number_format = num_fmt
                for cc in range(1, len(headers) + 1):
                    ws.cell(row=row, column=cc).border = border
                    if ci % 2 == 1 and not (eff_month is not None and 2 <= cc <= 13 and (cc - 2) >= eff_month):
                        ws.cell(row=row, column=cc).fill = PatternFill('solid', fgColor=COLORS['alt_row'])
                row += 1

            # Group subtotal
            for cc in range(1, len(headers) + 1):
                c = ws.cell(row=row, column=cc)
                c.fill = PatternFill('solid', fgColor='FFE0E7F0')
                c.font = Font(name='Arial', size=10, bold=True, color='1B2A4A')
                c.border = border
            ws.cell(row=row, column=1, value=f'Cəmi {group}')
            for col_idx in range(2, len(headers) + 1):
                cl = get_column_letter(col_idx)
                ws.cell(row=row, column=col_idx, value=f'=SUM({cl}{grp_start}:{cl}{row-1})')
                ws.cell(row=row, column=col_idx).number_format = num_fmt
            grand_total_rows.append(row)
            row += 1
            row += 1

        # Grand total
        for cc in range(1, len(headers) + 1):
            c = ws.cell(row=row, column=cc)
            c.fill = tot_fill
            c.font = tot_font
            c.border = border
        ws.cell(row=row, column=1, value='ÜMUMI CƏMİ')
        for col_idx in range(2, len(headers) + 1):
            cl = get_column_letter(col_idx)
            parts = [f'{cl}{r}' for r in grand_total_rows]
            ws.cell(row=row, column=col_idx, value=f'={"+".join(parts)}')
            ws.cell(row=row, column=col_idx).number_format = num_fmt

        # Bar chart
        chart_row = row + 2
        bar = BarChart()
        bar.type = 'col'
        bar.title = 'Aylıq ümumi satış dinamikası'
        bar.y_axis.title = 'AZN (₼)'
        bar.style = 10
        cats_ref = Reference(ws, min_col=2, max_col=13, min_row=3)
        vals_ref = Reference(ws, min_col=2, max_col=13, min_row=row)
        bar.add_data(vals_ref, from_rows=True)
        bar.set_categories(cats_ref)
        bar.shape = 4
        bar.width = 28
        bar.height = 14
        ws.add_chart(bar, f'B{chart_row}')

        ws.column_dimensions['A'].width = 40
        for ci in range(2, 15):
            ws.column_dimensions[get_column_letter(ci)].width = 16
        ws.sheet_view.showGridLines = False

    if view == 'service':
        ws = wb.create_sheet('Aylıq Satış (Xidmət)')
        ws.sheet_properties.tabColor = '2D7A4A'

        ws.merge_cells(f'A1:{get_column_letter(14)}1')
        title = f'Xidmət üzrə aylıq satış hesabatı — {now.year}'
        if effective_date:
            title += f'  (qiymət dəyişikliyi: {effective_date})'
        ws['A1'] = title
        ws['A1'].font = Font(name='Arial', size=14, bold=True, color='1B2A4A')

        # Aggregate services per-month using per-company effective dates
        # For each service, compute 12 monthly values by summing across companies
        # respecting each company's individual effective month
        def _build_service_monthly():
            """Returns {internal_cat: {svc_name: [month0..month11]}}"""
            result = {}
            for comp_name in set(list(base_data.keys()) + list(adjusted_data.keys())):
                base_info = base_data.get(comp_name, {})
                adj_info = adjusted_data.get(comp_name, {})
                group = (adj_info or base_info).get('group', '')

                for cat in set(list(base_info.get('categories', {}).keys()) + list(adj_info.get('categories', {}).keys())):
                    comp_eff = _get_company_eff_month(comp_name, group, adjustments, effective_date, category_name=cat)
                    base_cat = base_info.get('categories', {}).get(cat, {})
                    adj_cat = adj_info.get('categories', {}).get(cat, {})
                    if not isinstance(base_cat, dict):
                        base_cat = {'total': base_cat, 'services': []}
                    if not isinstance(adj_cat, dict):
                        adj_cat = {'total': adj_cat, 'services': []}

                    base_svcs = {s['name']: s.get('total', 0) or 0 for s in base_cat.get('services', [])}
                    adj_svcs = {s['name']: s.get('total', 0) or 0 for s in adj_cat.get('services', [])}

                    if cat not in result:
                        result[cat] = {}

                    for sn in set(list(base_svcs.keys()) + list(adj_svcs.keys())):
                        if sn not in result[cat]:
                            result[cat][sn] = [0.0] * 12
                        bv = base_svcs.get(sn, 0)
                        av = adj_svcs.get(sn, 0)
                        for mi in range(12):
                            use_adj = (comp_eff is None) or (mi >= comp_eff)
                            result[cat][sn][mi] += av if use_adj else bv
            return result

        svc_monthly = _build_service_monthly()

        headers = ['Xidmət'] + MONTHS_AZ + ['İllik Cəmi']
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=ci, value=h)
            c.font = hdr_font
            c.fill = hdr_fill
            c.border = border
            c.alignment = Alignment(horizontal='center', wrap_text=True)

        row = 4
        cat_total_rows = []

        for board_cat in BOARD_CATS:
            internal_cats = CATEGORY_MAP.get(board_cat, [])
            all_svc_names = set()
            for ic in internal_cats:
                all_svc_names.update(svc_monthly.get(ic, {}).keys())
            if not all_svc_names:
                continue

            svc_data = {}
            for ic in internal_cats:
                for sn, months in svc_monthly.get(ic, {}).items():
                    if sn in svc_data:
                        svc_data[sn] = [svc_data[sn][i] + months[i] for i in range(12)]
                    else:
                        svc_data[sn] = list(months)

            for cc in range(1, len(headers) + 1):
                c = ws.cell(row=row, column=cc)
                c.fill = grp_fill
                c.font = grp_font
                c.border = border
            ws.cell(row=row, column=1, value=f'▸ {board_cat}')
            row += 1
            cat_start = row

            for si, sname in enumerate(sorted(svc_data.keys())):
                ws.cell(row=row, column=1, value=sname).font = normal
                months_vals = svc_data[sname]
                has_diff = any(abs(months_vals[i] - months_vals[0]) > 0.01 for i in range(1, 12))
                for mi in range(12):
                    val = months_vals[mi]
                    # Show green if value differs from Jan (= adjusted months)
                    is_adj = has_diff and abs(val - months_vals[0]) > 0.01
                    c = ws.cell(row=row, column=2 + mi, value=val)
                    c.font = adj_font if is_adj else normal
                    c.number_format = num_fmt
                    if is_adj:
                        c.fill = adj_fill
                c = ws.cell(row=row, column=14, value=f'=SUM(B{row}:M{row})')
                c.font = bold
                c.number_format = num_fmt
                for cc in range(1, len(headers) + 1):
                    ws.cell(row=row, column=cc).border = border
                    if si % 2 == 1:
                        # Don't override green adjusted cells
                        existing = ws.cell(row=row, column=cc).fill
                        if not (existing and existing.fgColor and existing.fgColor.rgb and existing.fgColor.rgb != '00000000'):
                            ws.cell(row=row, column=cc).fill = PatternFill('solid', fgColor=COLORS['alt_row'])
                row += 1

            # Category subtotal
            for cc in range(1, len(headers) + 1):
                c = ws.cell(row=row, column=cc)
                c.fill = PatternFill('solid', fgColor='FFE0E7F0')
                c.font = Font(name='Arial', size=10, bold=True, color='1B2A4A')
                c.border = border
            ws.cell(row=row, column=1, value=f'Cəmi {board_cat}')
            for col_idx in range(2, len(headers) + 1):
                cl = get_column_letter(col_idx)
                ws.cell(row=row, column=col_idx, value=f'=SUM({cl}{cat_start}:{cl}{row-1})')
                ws.cell(row=row, column=col_idx).number_format = num_fmt
            cat_total_rows.append(row)
            row += 1
            row += 1

        # Grand total
        for cc in range(1, len(headers) + 1):
            c = ws.cell(row=row, column=cc)
            c.fill = tot_fill
            c.font = tot_font
            c.border = border
        ws.cell(row=row, column=1, value='ÜMUMI CƏMİ')
        for col_idx in range(2, len(headers) + 1):
            cl = get_column_letter(col_idx)
            parts = [f'{cl}{r}' for r in cat_total_rows]
            if parts:
                ws.cell(row=row, column=col_idx, value=f'={"+".join(parts)}')
            ws.cell(row=row, column=col_idx).number_format = num_fmt

        ws.column_dimensions['A'].width = 55
        for ci in range(2, 15):
            ws.column_dimensions[get_column_letter(ci)].width = 16
        ws.sheet_view.showGridLines = False

    return ws


def generate_template1(data, legal, adjustments=None, output_path=None, effective_date=None):
    base_data = data  # keep original for monthly sheet
    if adjustments:
        data = apply_adjustments(data, adjustments)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    num_fmt = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
    bold_font = Font(name='Calibri', size=11, bold=True)
    normal_font = Font(name='Calibri', size=11)

    headers = [''] + BOARD_CATS + ['Total']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = bold_font

    row = 2
    sorted_companies = sorted(data.keys(), key=lambda x: legal.get(x, x + ' MMC'))
    for name in sorted_companies:
        info = data[name]
        full_name = legal.get(name, name + ' MMC')
        ws.cell(row=row, column=1, value=full_name).font = bold_font
        board = aggregate_board_cats(info['categories'])
        for col_idx, cat in enumerate(BOARD_CATS, 2):
            c = ws.cell(row=row, column=col_idx, value=board.get(cat, 0))
            c.font = normal_font
            c.number_format = num_fmt
        total_col = len(BOARD_CATS) + 2
        c = ws.cell(row=row, column=total_col, value=f'=SUM(B{row}:{get_column_letter(total_col-1)}{row})')
        c.font = normal_font
        c.number_format = num_fmt
        row += 1

    ws.cell(row=row, column=1, value='Total ').font = bold_font
    for col_idx in range(2, len(BOARD_CATS) + 3):
        col_letter = get_column_letter(col_idx)
        c = ws.cell(row=row, column=col_idx, value=f'=SUM({col_letter}2:{col_letter}{row-1})')
        c.font = bold_font
        c.number_format = num_fmt

    ws.column_dimensions['A'].width = 50
    for col_idx in range(2, len(BOARD_CATS) + 3):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    # Add monthly sales sheets (base vs adjusted split by effective_date)
    add_monthly_sales_sheet(wb, base_data, data, legal, effective_date=effective_date, view='company', adjustments=adjustments)
    add_monthly_sales_sheet(wb, base_data, data, legal, effective_date=effective_date, view='service', adjustments=adjustments)

    if not output_path:
        output_path = os.path.join(STATIC_DIR, 'exports', 'SALES_Template1.xlsx')
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path


def generate_template2(data, legal, adjustments=None, output_path=None, effective_date=None):
    base_data = data  # keep original for monthly sheet
    if adjustments:
        data = apply_adjustments(data, adjustments)
    wb = Workbook()

    # --- Summary Sheet ---
    ws_sum = wb.active
    ws_sum.title = 'Xülasə'
    ws_sum.sheet_properties.tabColor = '1B2A4A'

    thin = Side(style='thin', color=COLORS['border'])
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    hdr_fill = PatternFill('solid', fgColor=COLORS['header_bg'])
    hdr_font = Font(name='Arial', size=11, bold=True, color=COLORS['header_font'])
    grp_fill = PatternFill('solid', fgColor=COLORS['group_bg'])
    grp_font = Font(name='Arial', size=11, bold=True, color=COLORS['group_font'])
    tot_fill = PatternFill('solid', fgColor=COLORS['total_bg'])
    tot_font = Font(name='Arial', size=12, bold=True, color=COLORS['total_font'])
    normal = Font(name='Arial', size=10)
    bold = Font(name='Arial', size=10, bold=True)
    num_fmt = '#,##0.00 [$₼-42C]'
    pct_fmt = '0.0%'

    title_font = Font(name='Arial', size=16, bold=True, color='1B2A4A')
    sub_font = Font(name='Arial', size=11, color='546E7A')
    now = datetime.datetime.now()
    ws_sum.merge_cells('A1:J1')
    ws_sum['A1'] = 'GUVEN TECHNOLOGY HOLDİNG'
    ws_sum['A1'].font = title_font
    ws_sum.merge_cells('A2:J2')
    ws_sum['A2'] = f'Aylıq Gəlir Hesabatı — {now.strftime("%B %Y")}'
    ws_sum['A2'].font = sub_font

    # KPI Row
    kpi_fill = PatternFill('solid', fgColor=COLORS['kpi_bg'])
    kpi_font = Font(name='Arial', size=20, bold=True, color=COLORS['kpi_val'])
    kpi_label = Font(name='Arial', size=9, color='546E7A')
    total_monthly = sum(d['monthly'] for d in data.values())
    total_annual = total_monthly * 12
    num_companies = len(data)
    avg_revenue = total_monthly / num_companies if num_companies else 0

    kpis = [
        ('Ümumi Aylıq Gəlir', total_monthly, num_fmt),
        ('Ümumi İllik Gəlir', total_annual, num_fmt),
        ('Şirkət Sayı', num_companies, '#,##0'),
        ('Orta Gəlir/Şirkət', avg_revenue, num_fmt),
    ]
    for i, (label, val, fmt) in enumerate(kpis):
        col = 1 + i * 3
        ws_sum.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col+1)
        c = ws_sum.cell(row=4, column=col, value=label)
        c.font = kpi_label
        ws_sum.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col+1)
        c = ws_sum.cell(row=5, column=col, value=val)
        c.font = kpi_font
        c.number_format = fmt
        for r in [4, 5]:
            for cc in range(col, col+2):
                ws_sum.cell(row=r, column=cc).fill = kpi_fill
                ws_sum.cell(row=r, column=cc).border = border

    # Group summary table
    row = 8
    grp_headers = ['Qrup', 'Şirkət Sayı', 'Aylıq Gəlir (₼)', 'İllik Gəlir (₼)', 'Pay (%)']
    for col_idx, h in enumerate(grp_headers, 1):
        c = ws_sum.cell(row=row, column=col_idx, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.border = border
        c.alignment = Alignment(horizontal='center')

    row = 9
    groups_data = {}
    for name, info in data.items():
        g = info['group']
        if g not in groups_data:
            groups_data[g] = {'count': 0, 'monthly': 0}
        groups_data[g]['count'] += 1
        groups_data[g]['monthly'] += info['monthly']

    grp_start = row
    for g in GROUP_ORDER:
        if g not in groups_data:
            continue
        gd = groups_data[g]
        ws_sum.cell(row=row, column=1, value=g).font = bold
        ws_sum.cell(row=row, column=2, value=gd['count']).font = normal
        ws_sum.cell(row=row, column=2).alignment = Alignment(horizontal='center')
        c = ws_sum.cell(row=row, column=3, value=gd['monthly'])
        c.font = normal
        c.number_format = num_fmt
        c = ws_sum.cell(row=row, column=4, value=gd['monthly'] * 12)
        c.font = normal
        c.number_format = num_fmt
        c = ws_sum.cell(row=row, column=5, value=gd['monthly'] / total_monthly if total_monthly else 0)
        c.font = normal
        c.number_format = pct_fmt
        for cc in range(1, 6):
            ws_sum.cell(row=row, column=cc).border = border
            if (row - grp_start) % 2 == 1:
                ws_sum.cell(row=row, column=cc).fill = PatternFill('solid', fgColor=COLORS['alt_row'])
        row += 1

    # Total row
    for cc in range(1, 6):
        ws_sum.cell(row=row, column=cc).fill = tot_fill
        ws_sum.cell(row=row, column=cc).font = tot_font
        ws_sum.cell(row=row, column=cc).border = border
    ws_sum.cell(row=row, column=1, value='CƏMİ')
    ws_sum.cell(row=row, column=2, value=f'=SUM(B{grp_start}:B{row-1})')
    ws_sum.cell(row=row, column=2).alignment = Alignment(horizontal='center')
    ws_sum.cell(row=row, column=3, value=f'=SUM(C{grp_start}:C{row-1})')
    ws_sum.cell(row=row, column=3).number_format = num_fmt
    ws_sum.cell(row=row, column=4, value=f'=SUM(D{grp_start}:D{row-1})')
    ws_sum.cell(row=row, column=4).number_format = num_fmt
    ws_sum.cell(row=row, column=5, value=1)
    ws_sum.cell(row=row, column=5).number_format = pct_fmt

    grp_end = row

    # Pie chart for group distribution
    pie = PieChart()
    pie.title = 'Qrup üzrə gəlir paylanması'
    pie.style = 10
    cats_ref = Reference(ws_sum, min_col=1, min_row=grp_start, max_row=grp_end - 1)
    vals_ref = Reference(ws_sum, min_col=3, min_row=grp_start, max_row=grp_end - 1)
    pie.add_data(vals_ref)
    pie.set_categories(cats_ref)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showCatName = True
    pie.width = 18
    pie.height = 12
    ws_sum.add_chart(pie, f'G8')

    ws_sum.column_dimensions['A'].width = 20
    ws_sum.column_dimensions['B'].width = 14
    ws_sum.column_dimensions['C'].width = 22
    ws_sum.column_dimensions['D'].width = 22
    ws_sum.column_dimensions['E'].width = 12

    # --- Category Summary Sheet ---
    ws_cat = wb.create_sheet('Kateqoriyalar')
    ws_cat.sheet_properties.tabColor = '2D4A7A'

    ws_cat.merge_cells('A1:I1')
    ws_cat['A1'] = 'Kateqoriya üzrə gəlir bölgüsü'
    ws_cat['A1'].font = Font(name='Arial', size=14, bold=True, color='1B2A4A')

    cat_headers = ['Kateqoriya', 'Aylıq Gəlir (₼)', 'İllik Gəlir (₼)', 'Pay (%)']
    for ci, h in enumerate(cat_headers, 1):
        c = ws_cat.cell(row=3, column=ci, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.border = border
        c.alignment = Alignment(horizontal='center')

    cat_totals = {bc: 0 for bc in BOARD_CATS}
    for info in data.values():
        board = aggregate_board_cats(info['categories'])
        for bc, val in board.items():
            cat_totals[bc] += val

    cat_start = 4
    for ri, bc in enumerate(BOARD_CATS):
        r = cat_start + ri
        ws_cat.cell(row=r, column=1, value=bc).font = bold
        c = ws_cat.cell(row=r, column=2, value=cat_totals[bc])
        c.font = normal
        c.number_format = num_fmt
        c = ws_cat.cell(row=r, column=3, value=cat_totals[bc] * 12)
        c.font = normal
        c.number_format = num_fmt
        c = ws_cat.cell(row=r, column=4, value=cat_totals[bc] / total_monthly if total_monthly else 0)
        c.font = normal
        c.number_format = pct_fmt
        for cc in range(1, 5):
            ws_cat.cell(row=r, column=cc).border = border
            if ri % 2 == 1:
                ws_cat.cell(row=r, column=cc).fill = PatternFill('solid', fgColor=COLORS['alt_row'])

    cat_end_row = cat_start + len(BOARD_CATS)
    for cc in range(1, 5):
        ws_cat.cell(row=cat_end_row, column=cc).fill = tot_fill
        ws_cat.cell(row=cat_end_row, column=cc).font = tot_font
        ws_cat.cell(row=cat_end_row, column=cc).border = border
    ws_cat.cell(row=cat_end_row, column=1, value='CƏMİ')
    ws_cat.cell(row=cat_end_row, column=2, value=f'=SUM(B{cat_start}:B{cat_end_row-1})')
    ws_cat.cell(row=cat_end_row, column=2).number_format = num_fmt
    ws_cat.cell(row=cat_end_row, column=3, value=f'=SUM(C{cat_start}:C{cat_end_row-1})')
    ws_cat.cell(row=cat_end_row, column=3).number_format = num_fmt
    ws_cat.cell(row=cat_end_row, column=4, value=1)
    ws_cat.cell(row=cat_end_row, column=4).number_format = pct_fmt

    bar = BarChart()
    bar.type = 'col'
    bar.title = 'Kateqoriya üzrə aylıq gəlir'
    bar.y_axis.title = 'AZN (₼)'
    bar.style = 10
    cats_ref = Reference(ws_cat, min_col=1, min_row=cat_start, max_row=cat_end_row - 1)
    vals_ref = Reference(ws_cat, min_col=2, min_row=cat_start, max_row=cat_end_row - 1)
    bar.add_data(vals_ref)
    bar.set_categories(cats_ref)
    bar.shape = 4
    bar.width = 22
    bar.height = 14
    ws_cat.add_chart(bar, 'F3')

    ws_cat.column_dimensions['A'].width = 25
    ws_cat.column_dimensions['B'].width = 22
    ws_cat.column_dimensions['C'].width = 22
    ws_cat.column_dimensions['D'].width = 12

    # --- Detail Sheet (same as Template 1 but with group subtotals and formatting) ---
    ws_det = wb.create_sheet('Ətraflı')
    ws_det.sheet_properties.tabColor = '4A7A2D'

    ws_det.merge_cells('A1:I1')
    ws_det['A1'] = 'Şirkət üzrə ətraflı gəlir hesabatı'
    ws_det['A1'].font = Font(name='Arial', size=14, bold=True, color='1B2A4A')

    det_headers = ['Şirkət'] + BOARD_CATS + ['Cəmi']
    for ci, h in enumerate(det_headers, 1):
        c = ws_det.cell(row=3, column=ci, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.border = border
        c.alignment = Alignment(horizontal='center', wrap_text=True)

    ws_det.row_dimensions[3].height = 30

    row = 4
    grand_total_rows = []
    for gi, group in enumerate(GROUP_ORDER):
        companies = [(n, d) for n, d in data.items() if d['group'] == group]
        if not companies:
            continue
        companies.sort(key=lambda x: x[1]['monthly'], reverse=True)

        for cc in range(1, len(det_headers) + 1):
            c = ws_det.cell(row=row, column=cc)
            c.fill = grp_fill
            c.font = grp_font
            c.border = border
        ws_det.cell(row=row, column=1, value=f'▸ {group}')
        row += 1
        grp_start_row = row

        for ci, (name, info) in enumerate(companies):
            full_name = legal.get(name, name + ' MMC')
            ws_det.cell(row=row, column=1, value=full_name).font = normal
            board = aggregate_board_cats(info['categories'])
            for col_idx, cat in enumerate(BOARD_CATS, 2):
                c = ws_det.cell(row=row, column=col_idx, value=board.get(cat, 0))
                c.font = normal
                c.number_format = num_fmt
            total_col = len(BOARD_CATS) + 2
            c = ws_det.cell(row=row, column=total_col, value=f'=SUM(B{row}:{get_column_letter(total_col-1)}{row})')
            c.font = normal
            c.number_format = num_fmt
            for cc in range(1, len(det_headers) + 1):
                ws_det.cell(row=row, column=cc).border = border
                if ci % 2 == 1:
                    ws_det.cell(row=row, column=cc).fill = PatternFill('solid', fgColor=COLORS['alt_row'])
            row += 1

        # Group subtotal
        for cc in range(1, len(det_headers) + 1):
            c = ws_det.cell(row=row, column=cc)
            c.fill = PatternFill('solid', fgColor='FFE0E7F0')
            c.font = Font(name='Arial', size=10, bold=True, color='1B2A4A')
            c.border = border
        ws_det.cell(row=row, column=1, value=f'Cəmi {group}')
        for col_idx in range(2, len(det_headers) + 1):
            cl = get_column_letter(col_idx)
            ws_det.cell(row=row, column=col_idx, value=f'=SUM({cl}{grp_start_row}:{cl}{row-1})')
            ws_det.cell(row=row, column=col_idx).number_format = num_fmt
        grand_total_rows.append(row)
        row += 1
        row += 1  # empty row between groups

    # Grand total
    for cc in range(1, len(det_headers) + 1):
        c = ws_det.cell(row=row, column=cc)
        c.fill = tot_fill
        c.font = tot_font
        c.border = border
    ws_det.cell(row=row, column=1, value='ÜMUMI CƏMİ')
    for col_idx in range(2, len(det_headers) + 1):
        cl = get_column_letter(col_idx)
        formula_parts = [f'{cl}{r}' for r in grand_total_rows]
        ws_det.cell(row=row, column=col_idx, value=f'={"+".join(formula_parts)}')
        ws_det.cell(row=row, column=col_idx).number_format = num_fmt

    ws_det.column_dimensions['A'].width = 50
    for col_idx in range(2, len(det_headers) + 1):
        ws_det.column_dimensions[get_column_letter(col_idx)].width = 18

    ws_det.sheet_view.showGridLines = False

    # Add monthly sales sheets (base vs adjusted split by effective_date)
    add_monthly_sales_sheet(wb, base_data, data, legal, effective_date=effective_date, view='company', adjustments=adjustments)
    add_monthly_sales_sheet(wb, base_data, data, legal, effective_date=effective_date, view='service', adjustments=adjustments)

    if not output_path:
        output_path = os.path.join(STATIC_DIR, 'exports', 'SALES_Template2.xlsx')
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path


BUDGET_CAT_MAP = {
    'Sales revenue': BOARD_CATS,  # all categories summed = total revenue
    'Permanent IT service': ['Daimi İT xidməti'],
    'ERP': ['ERP'],
    'Additional IT services': ['Əlavə IT xidməti'],
    'Information security services': ['İnfosec'],
    'GRS': ['GRC'],
    'Projects': ['Layihə'],
}

BUDGET_ROWS_ORDER = [
    'Sales revenue',
    'Permanent IT service',
    'ERP',
    'Additional IT services',
    'Information security services',
    'GRS',
    'Projects',
]


def generate_budget_pl(data, legal, adjustments=None, output_path=None, effective_date=None):
    """Generate Budget P&L template: categories as rows, months as columns."""
    base_data = data
    if adjustments:
        data = apply_adjustments(data, adjustments)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Budget PL'

    now = datetime.datetime.now()
    year = now.year

    # Styles matching the uploaded template
    title_font = Font(name='Calibri', size=14, bold=True)
    hdr_font = Font(name='Calibri', size=11, bold=True)
    normal_font = Font(name='Calibri', size=11)
    revenue_fmt = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
    detail_fmt = '#,##0.00'
    annual_fmt = '#,##0'
    date_fmt = '[$-409]mmm\\-yy;@'

    thin = Side(style='thin', color='FFB0BEC5')
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    # Row 1: Title
    ws['A1'] = f'BUDGET {year}'
    ws['A1'].font = title_font
    ws['O1'] = 'Budget'
    ws['O1'].font = hdr_font

    # Row 2: Month headers (dates for each month)
    for mi in range(12):
        # Use ~end of month dates like in template (last business day area)
        import calendar
        last_day = calendar.monthrange(year, mi + 1)[1]
        dt = datetime.datetime(year, mi + 1, last_day)
        c = ws.cell(row=2, column=3 + mi, value=dt)
        c.number_format = date_fmt
        c.font = hdr_font
        c.alignment = Alignment(horizontal='center')
        c.border = border
    ws.cell(row=2, column=15, value=str(year))
    ws.cell(row=2, column=15).font = hdr_font
    ws.cell(row=2, column=15).alignment = Alignment(horizontal='center')

    # Build monthly totals per board category
    # For each company: base months use base_data, adjusted months use adjusted data
    cat_monthly = {bc: [0.0] * 12 for bc in BOARD_CATS}

    for comp_name in set(list(base_data.keys()) + list(data.keys())):
        base_info = base_data.get(comp_name, {})
        adj_info = data.get(comp_name, {})
        group = (adj_info or base_info).get('group', '')

        base_board = aggregate_board_cats(base_info.get('categories', {}))
        adj_board = aggregate_board_cats(adj_info.get('categories', {}))

        eff_month = _get_company_eff_month(comp_name, group, adjustments, effective_date)

        for bc in BOARD_CATS:
            bv = base_board.get(bc, 0)
            av = adj_board.get(bc, 0)
            for mi in range(12):
                use_adj = (eff_month is None) or (mi >= eff_month)
                cat_monthly[bc][mi] += av if use_adj else bv

    # Write data rows
    row = 3
    for cat_label in BUDGET_ROWS_ORDER:
        ws.cell(row=row, column=1, value=cat_label).font = normal_font

        board_cats_for_row = BUDGET_CAT_MAP[cat_label]
        monthly_vals = [0.0] * 12
        for bc in board_cats_for_row:
            for mi in range(12):
                monthly_vals[mi] += cat_monthly.get(bc, [0.0] * 12)[mi]

        is_total_row = (cat_label == 'Sales revenue')
        fmt = revenue_fmt if is_total_row else detail_fmt

        for mi in range(12):
            c = ws.cell(row=row, column=3 + mi, value=round(monthly_vals[mi], 2))
            c.number_format = fmt
            c.font = Font(name='Calibri', size=11, bold=is_total_row)
            c.border = border

        # Annual total (column O = 15)
        annual = sum(monthly_vals)
        c = ws.cell(row=row, column=15, value=round(annual, 2))
        c.number_format = annual_fmt
        c.font = Font(name='Calibri', size=11, bold=is_total_row)
        c.border = border

        row += 1

    # Column widths
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 4
    for ci in range(3, 16):
        ws.column_dimensions[get_column_letter(ci)].width = 14

    if not output_path:
        output_path = os.path.join(STATIC_DIR, 'exports', 'Budget_PL.xlsx')
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == '__main__':
    data, legal = load_data()
    p1 = generate_template1(data, legal)
    print(f'Template 1: {p1}')
    p2 = generate_template2(data, legal)
    print(f'Template 2: {p2}')
    p3 = generate_budget_pl(data, legal)
    print(f'Budget PL: {p3}')
